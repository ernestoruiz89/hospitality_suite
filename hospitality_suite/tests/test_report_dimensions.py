"""Native transactions exercise report scope, history and financial sources."""
from datetime import datetime, timedelta
import io
import uuid

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import now_datetime
from openpyxl import load_workbook

from hospitality_suite import api
from hospitality_suite.pricing import decimal
from hospitality_suite.service import operation
from hospitality_suite.time import to_site


class TestReportDimensions(IntegrationTestCase):
    def setUp(self):
        super().setUp()
        self.addCleanup(frappe.db.rollback)
        frappe.set_user("Administrator")
        self.prop=frappe.get_doc("HS Property",{"property_name":"Hotel Jardín Demo"})
        self.rooms=frappe.get_all("HS Room",filters={"property":self.prop.name},fields=["name","category","room_number"],order_by="room_number")
        self.room=self.rooms[0]
        self.guest=frappe.get_all("HS Guest",filters={"property":self.prop.name},fields=["name","customer"])[0]

    def tearDown(self):
        frappe.set_user("Administrator")
        super().tearDown()

    def call(self,action,data):
        return api.execute(self.prop.name,action,uuid.uuid4().hex,data)

    def enter(self,room=None,people=1):
        room=room or self.room
        start=now_datetime().replace(microsecond=0)
        values=dict(starts_at=str(start),ends_at=str(start+timedelta(days=1)),guest=self.guest.name,customer=self.guest.customer,
                    channel="Agencia de prueba",rooms=[dict(category=room.category,room=room.name,people=people)])
        values["acceptance"]=api.preview(self.prop.name,values)["acceptance"]
        return self.call("walk_in",values)

    def test_arrival_departure_and_registration_are_distinct(self):
        start=now_datetime().replace(microsecond=0)+timedelta(days=15)
        end=start+timedelta(days=2)
        values=dict(starts_at=str(start),ends_at=str(end),customer=self.guest.customer,guest=self.guest.name,
                    channel="Agencia de prueba",rooms=[dict(category=self.room.category,room=self.room.name,people=1)])
        values["acceptance"]=api.preview(self.prop.name,values)["acceptance"]
        booking=self.call("reserve",values)["booking"]
        selected=dict(from_date=str(start.date()),to_date=str(start.date()),room=self.room.name,category=self.room.category,channel="Agencia de prueba",user="Administrator")
        self.assertEqual([r.name for r in api.report(self.prop.name,"Llegadas",selected)],[booking])
        self.assertFalse(api.report(self.prop.name,"Reservas",dict(selected,date_field="creation")))
        self.assertEqual([r.name for r in api.report(self.prop.name,"Reservas",dict(selected,date_field="ends_at",from_date=str(end.date()),to_date=str(end.date())))],[booking])
        self.assertFalse(api.report(self.prop.name,"Reservas",dict(selected,channel="Otro canal")))
        cancelled=frappe.get_doc("HS Booking",booking)
        self.call("booking_status",dict(booking=booking,status="Cancelled",reason="Cancelación sintética",modified=str(cancelled.modified)))
        self.assertEqual([r.name for r in api.report(self.prop.name,"Cancelaciones y no presentaciones",dict(selected,date_field="starts_at"))],[booking])
        self.assertFalse(api.report(self.prop.name,"Llegadas",selected))

    def test_local_day_boundaries_one_sided_period_and_display(self):
        with operation():
            self.prop.time_zone="Asia/Tokyo"
            self.prop.save()
        start=to_site("2027-01-02 00:00:00",self.prop)
        end=to_site("2027-01-03 00:00:00",self.prop)
        prefix="period-"+uuid.uuid4().hex[:8]
        stamps=[start-timedelta(microseconds=1),start,end-timedelta(microseconds=1),end]
        fields=["name","company","property","action","actor","occurred_at","owner","creation","modified"]
        frappe.db.bulk_insert("HS Audit Event",fields=fields,values=[(prefix+str(i),self.prop.company,self.prop.name,"Period fixture","Administrator",stamp,"Administrator",stamp,stamp) for i,stamp in enumerate(stamps)])
        result=api.report(self.prop.name,"Auditoría",dict(from_date="2027-01-02",to_date="2027-01-02",user="Administrator"))
        self.assertEqual({r.name for r in result},{prefix+"1",prefix+"2"})
        self.assertTrue(all("+09:00" in r._display["occurred_at"] for r in result))
        one_sided=api.report(self.prop.name,"Auditoría",dict(from_date="2027-01-03"))
        self.assertEqual({r.name for r in one_sided},{prefix+"3"})
        page=api.report_page(self.prop.name,"Auditoría",dict(from_date="2027-01-02",to_date="2027-01-02"))
        self.assertNotIn("_display",page["columns"])
        for bad in (dict(room=self.room.name),dict(from_date="2027-02-30"),dict(from_date="2027-02-02",to_date="2027-01-01"),dict(date_field="unknown")):
            with self.assertRaises(frappe.ValidationError):
                api.report(self.prop.name,"Auditoría",bad)

    def test_hosted_guests_include_principal_and_companions_with_names(self):
        entered=self.enter(people=2)
        stay=frappe.get_doc("HS Stay",entered["stay"])
        companion=frappe.get_all("HS Guest",filters={"property":self.prop.name,"name":["!=",self.guest.name]},fields=["name","full_name"])[0]
        self.call("companions",dict(stay=stay.name,modified=str(stay.modified),guests=[companion.name]))
        rows=api.report(self.prop.name,"Huéspedes alojados",dict(room=self.room.name))
        self.assertEqual({r.guest:r.relationship for r in rows},{self.guest.name:"Principal",companion.name:"Acompañante"})
        self.assertEqual(next(r._display["guest"] for r in rows if r.guest==companion.name),companion.full_name)
        self.assertTrue(all(r.name==stay.name for r in rows))

    def test_production_room_and_adjustment_keep_original_attribution_after_move(self):
        entered=self.enter()
        charge=frappe.db.get_value("HS Charge",{"stay":entered["stay"],"kind":"Lodging"},"name")
        stay=frappe.get_doc("HS Stay",entered["stay"])
        values=dict(stay=stay.name,room=self.rooms[1].name,folio=entered["folio"],reason="Cambio sintético",modified=str(stay.modified))
        values["acceptance"]=api.preview(self.prop.name,values,"change_room")["acceptance"]
        self.call("change_room",values)
        adjustment=self.call("adjustment",dict(charge=charge,adjustment_type="Discount",amount=100,reason="Descuento posterior al cambio"))["charge"]
        results=api.report(self.prop.name,"Producción",dict(room=self.room.name,category=self.room.category,channel="Agencia de prueba"))
        self.assertTrue({charge,adjustment}.issubset({r.name for r in results}))
        self.assertFalse({charge,adjustment}.intersection(r.name for r in api.report(self.prop.name,"Producción",dict(room=self.rooms[1].name))))
        discounts=api.report(self.prop.name,"Descuentos y cortesías",dict(room=self.room.name,adjustment_type="Discount"))
        self.assertEqual([r.name for r in discounts],[adjustment])
        self.assertEqual(discounts[0]._display["room"],self.room.room_number)

    def test_invoice_credit_and_discount_reconcile_production_and_typed_export(self):
        entered=self.enter()
        service=self.call("request_service",dict(stay=entered["stay"],folio=entered["folio"],item="HS-DEMO-WATER",quantity=2))["service"]
        requested=frappe.get_doc("HS Service Request",service)
        self.call("update_service",dict(service=service,status="Delivered",modified=str(requested.modified)))
        charge=frappe.db.get_value("HS Service Request",service,"charge")
        self.call("adjustment",dict(charge=charge,adjustment_type="Discount",amount=6,reason="Descuento sintético"))
        invoice=self.call("invoice",dict(folio=entered["folio"]))["invoice"]
        inv=frappe.get_doc("Sales Invoice",invoice)
        line=next(r for r in inv.items if r.item_code=="HS-DEMO-WATER")
        credit=self.call("credit_note",dict(invoice=invoice,reason="Devolución parcial sintética",items=[dict(invoice_item=line.name,quantity=1)]))
        rows=api.report(self.prop.name,"Producción",dict(room=self.room.name,item="HS-DEMO-WATER"))
        self.assertEqual(len(rows),3)
        self.assertEqual(sum(decimal(r.amount) for r in rows),decimal(27))
        self.assertEqual(sum(decimal(r.revenue_net) for r in rows),decimal(27))
        returned=[r for r in rows if r.sales_invoice==credit["credit_note"]]
        self.assertEqual(returned[0].source_type,"Nota de crédito")
        self.assertEqual(returned[0]._display["room"],self.room.room_number)
        grouped=api.report(self.prop.name,"Producción",dict(room=self.room.name,item="HS-DEMO-WATER",group_by="item"))
        self.assertEqual(len(grouped),1)
        self.assertEqual((grouped[0].amount,grouped[0].revenue_net,grouped[0].pending_amount,grouped[0].records),(decimal(27),decimal(27),decimal(0),3))
        frappe.local.response=frappe._dict()
        api.export_report(self.prop.name,"Producción",dict(room=self.room.name,item="HS-DEMO-WATER"),"xlsx")
        wb=load_workbook(io.BytesIO(frappe.local.response.filecontent),read_only=True)
        values=list(wb.active.values)
        self.assertNotIn("_display",values[0])
        amount=values[0].index("Importe")
        room=values[0].index("Habitación")
        self.assertTrue(all(isinstance(row[amount],(int,float)) and row[room]==self.room.room_number for row in values[1:]))
        wb.close()

    def test_direct_sale_return_is_in_production_once_and_filterable(self):
        shift=self.call("open_shift",dict(opening_amount=0))["shift"]
        method=frappe.db.get_value("HS Payment Method",{"property":self.prop.name})
        values=dict(customer=self.guest.customer,item="HS-DEMO-WATER",quantity=2,method=method,shift=shift)
        values["acceptance"]=api.preview(self.prop.name,values,"direct_sale")["acceptance"]
        sale=self.call("direct_sale",values)
        self.call("return_direct_sale",dict(service=sale["service"],shift=shift,reason="Retorno de prueba"))
        rows=api.report(self.prop.name,"Producción",dict(channel="Venta directa",item="HS-DEMO-WATER",user="Administrator"))
        self.assertEqual(len(rows),2)
        self.assertEqual(sum(decimal(r.amount) for r in rows),0)
        self.assertEqual(sum(decimal(r.revenue_net) for r in rows),0)
        self.assertEqual({r.source_type for r in rows},{"Venta directa","Devolución de venta directa"})
        payments=api.report(self.prop.name,"Cobros",dict(shift=shift,method=method,user="Administrator"))
        self.assertEqual({r.kind for r in payments},{"Payment","Refund"})
        self.assertFalse(api.report(self.prop.name,"Cobros",dict(shift="other-shift")))

    def test_booking_penalties_use_decision_date_and_native_financial_sources(self):
        from hospitality_suite.time import local
        from hospitality_suite import reports
        for status,days,amount in (("Cancelled",15,100),("No Show",-2,80)):
            with self.subTest(status=status):
                with operation():
                    self.prop.cancellation_fee_mode=self.prop.no_show_fee_mode="Fixed"
                    self.prop.cancellation_fee_value=100
                    self.prop.no_show_fee_value=80
                    self.prop.penalty_item=self.prop.lodging_item
                    self.prop.save()
                start=now_datetime().replace(microsecond=0)+timedelta(days=days)
                channel="Penalización "+status
                values=dict(starts_at=str(start),ends_at=str(start+timedelta(days=1)),customer=self.guest.customer,guest=self.guest.name,
                            channel=channel,rooms=[dict(category=self.room.category,room=self.room.name,people=1)])
                values["acceptance"]=api.preview(self.prop.name,values)["acceptance"]
                booking=self.call("reserve",values)["booking"]
                doc=frappe.get_doc("HS Booking",booking)
                decision=dict(booking=booking,modified=str(doc.modified),status=status,reason="Decisión sintética")
                decision["fee_acceptance"]=api.preview(self.prop.name,decision,"booking_status")["acceptance"]
                self.call("booking_status",decision)
                doc.reload()
                invoice=frappe.get_doc("Sales Invoice",doc.fee_invoice)
                self.assertFalse(frappe.db.exists("HS Charge",{"sales_invoice":invoice.name}))
                today=str(local(now_datetime(),self.prop).date())
                selected=dict(from_date=today,to_date=today,status=status,channel=channel)
                result=api.report(self.prop.name,"Cancelaciones y no presentaciones",selected)
                self.assertEqual([r.name for r in result],[booking])
                self.assertEqual(result[0].penalty_amount,amount)
                self.assertEqual(result[0].fee_invoice,invoice.name)
                self.assertEqual(result[0].fee_status,"Confirmada")
                self.assertFalse(api.report(self.prop.name,"Cancelaciones y no presentaciones",dict(selected,from_date=str(start.date()),to_date=str(start.date()))))
                self.assertEqual([r.name for r in api.report(self.prop.name,"Cancelaciones y no presentaciones",dict(selected,date_field="starts_at",from_date=str(start.date()),to_date=str(start.date())))],[booking])
                frappe.db.set_value("HS Booking",booking,"modified",now_datetime()+timedelta(days=10),update_modified=False)
                self.assertEqual(api.report(self.prop.name,"Cancelaciones y no presentaciones",selected)[0].status_changed_at,result[0].status_changed_at)
                filters=dict(channel=channel,room=self.room.name,item=self.prop.penalty_item)
                production=api.report(self.prop.name,"Producción",filters)
                self.assertEqual(len(production),1)
                self.assertEqual((decimal(production[0].amount),decimal(production[0].revenue_net)),(decimal(invoice.total),decimal(invoice.net_total)))
                self.assertNotIn("Venta directa",production[0].channel)
                self.assertTrue(production[0].source_type.startswith("Penalización"))
                self.assertEqual(production[0].business_date,self.prop.business_date)
                view=api.report_page(self.prop.name,"Producción",filters)
                self.assertIn("Penalización",view["rows"][0]._display["item"])
                before=decimal(reports.summary(self.prop)["production"])
                credit=self.call("credit_note",dict(invoice=invoice.name,reason="Devolución de penalización",items=[dict(invoice_item=invoice.items[0].name,quantity=1)]))
                rows=api.report(self.prop.name,"Producción",filters)
                self.assertEqual(len(rows),2)
                self.assertEqual(sum(decimal(r.amount) for r in rows),0)
                self.assertEqual(sum(decimal(r.revenue_net) for r in rows),0)
                self.assertEqual(decimal(reports.summary(self.prop)["production"]),before-decimal(amount))
                self.assertEqual(next(r for r in rows if r.sales_invoice==credit["credit_note"]).source_type,"Devolución · "+production[0].source_type)
                for name in (credit["credit_note"],invoice.name):
                    native=frappe.get_doc("Sales Invoice",name)
                    native.hs_correction_reason="Anulación sintética autorizada"
                    native.save()
                    native.cancel()
                self.assertFalse(api.report(self.prop.name,"Producción",filters))
                self.assertEqual(api.report(self.prop.name,"Cancelaciones y no presentaciones",selected)[0].fee_status,"Anulada")

    def test_native_counter_invoice_and_credit_are_in_production_without_stay(self):
        invoice=frappe.get_doc(dict(doctype="Sales Invoice",company=self.prop.company,hs_property=self.prop.name,
                                    customer=self.guest.customer,currency=self.prop.currency,debit_to=self.prop.receivable_account,
                                    cost_center=self.prop.cost_center,update_stock=1))
        invoice.append("items",dict(item_code="HS-DEMO-WATER",qty=2,rate=30,cost_center=self.prop.cost_center,warehouse=self.prop.warehouse))
        invoice.set_missing_values()
        invoice.insert()
        invoice.submit()
        self.assertFalse(frappe.db.exists("HS Service Request",{"sales_invoice":invoice.name}))
        self.assertFalse(frappe.db.exists("HS Charge",{"sales_invoice":invoice.name}))
        rows=api.report(self.prop.name,"Producción",dict(channel="Venta directa",item="HS-DEMO-WATER"))
        original=[r for r in rows if r.sales_invoice==invoice.name]
        self.assertEqual(len(original),1)
        self.assertEqual((decimal(original[0].amount),decimal(original[0].revenue_net)),(decimal(60),decimal(60)))
        credit=self.call("credit_note",dict(invoice=invoice.name,reason="Devolución nativa parcial",items=[dict(invoice_item=invoice.items[0].name,quantity=1)]))
        rows=api.report(self.prop.name,"Producción",dict(channel="Venta directa",item="HS-DEMO-WATER"))
        chosen=[r for r in rows if r.sales_invoice in (invoice.name,credit["credit_note"])]
        self.assertEqual(len(chosen),2)
        self.assertEqual(sum(decimal(r.amount) for r in chosen),decimal(30))
        self.assertEqual(sum(decimal(r.revenue_net) for r in chosen),decimal(30))
        returned=frappe.get_doc("Sales Invoice",credit["credit_note"])
        self.assertTrue(returned.update_stock)
        stock=frappe.get_all("Stock Ledger Entry",filters={"voucher_type":"Sales Invoice","voucher_no":["in",[invoice.name,returned.name]],"is_cancelled":0},fields=["voucher_no","actual_qty"])
        self.assertEqual(sum(decimal(r.actual_qty) for r in stock if r.voucher_no==invoice.name),-2)
        self.assertEqual(sum(decimal(r.actual_qty) for r in stock if r.voucher_no==returned.name),1)
        from hospitality_suite import reports
        before=decimal(reports.summary(self.prop)["production"])
        foreign=frappe.get_doc(dict(doctype="Sales Invoice",company=self.prop.company,hs_property=self.prop.name,
                                    customer=self.guest.customer,currency="USD",conversion_rate=37,debit_to=self.prop.receivable_account,
                                    cost_center=self.prop.cost_center,update_stock=0))
        foreign.append("items",dict(item_code=self.prop.lodging_item,qty=1,rate=5,cost_center=self.prop.cost_center,warehouse=self.prop.warehouse))
        foreign.set_missing_values()
        foreign.insert()
        foreign.submit()
        totals=reports.summary(self.prop)
        self.assertEqual(decimal(totals["production"]),before,"Do not add USD directly to NIO")
        self.assertEqual(decimal(totals["production_by_currency"]["USD"]),5)
        self.assertEqual(decimal(totals["production_by_currency"][self.prop.currency]),before)

    def test_cleaning_productivity_and_shift_counts_preserve_their_units(self):
        from hospitality_suite.service import insert
        start=now_datetime()-timedelta(hours=3)
        with operation():
            for minutes in (30,60):
                insert("HS Cleaning Task",self.prop,room=self.room.name,kind="Deep",priority="Normal",status="Completed",assigned_to="Administrator",
                       started_at=start,completed_at=start+timedelta(minutes=minutes))
            insert("HS Cleaning Task",self.prop,room=self.room.name,kind="Stay",priority="Normal",status="Pending",assigned_to="Administrator")
            insert("HS Cleaning Task",self.prop,room=self.room.name,kind="Deep",priority="Normal",status="Completed")
        records=api.report(self.prop.name,"Productividad de limpieza",dict(room=self.room.name,assigned_to="Administrator"))
        self.assertEqual(len(records),1)
        result=records[0]
        self.assertEqual((result.tasks,result.pending_tasks,result.completed_tasks,result.timed_tasks),(3,1,2,2))
        self.assertEqual((result.duration_minutes,result.average_minutes,result.tasks_per_recorded_hour),(90,45,1.33))
        missing=next(r for r in api.report(self.prop.name,"Productividad de limpieza",dict(room=self.room.name)) if not r.assigned_to)
        self.assertEqual((missing.completed_tasks,missing.timed_tasks,missing.tasks_per_recorded_hour),(1,0,0))
        page=api.report_page(self.prop.name,"Productividad de limpieza",dict(room=self.room.name,assigned_to="Administrator"))
        self.assertEqual(page["rows"][0]._display["assigned_to"],frappe.db.get_value("User","Administrator","full_name"))
        shift=self.call("open_shift",dict(opening_amount=100))["shift"]
        before=api.report(self.prop.name,"Arqueos por moneda",dict(shift=shift))
        self.assertIsNone(before[0].counted_amount)
        self.assertIsNone(before[0].difference)
        doc=frappe.get_doc("HS Shift",shift)
        self.call("close_shift",dict(shift=shift,modified=str(doc.modified),denominations={"NIO":[dict(value=100,count=1)],"USD":[dict(value=1,count=5)]},
                                     reason="Sobrante USD sintético",approve=True))
        rows=api.report(self.prop.name,"Arqueos por moneda",dict(shift=shift,status="Closed"))
        currencies={r.currency:r for r in rows}
        self.assertEqual(set(currencies),{"NIO","USD"})
        self.assertEqual((currencies["NIO"].expected_amount,currencies["NIO"].counted_amount,currencies["NIO"].difference),(100,100,0))
        self.assertEqual((currencies["USD"].expected_amount,currencies["USD"].counted_amount,currencies["USD"].difference),(0,5,5))
        self.assertEqual(frappe.db.get_value("HS Shift",shift,"approved_by"),"Administrator")
        self.assertFalse(frappe.db.exists("Journal Entry",{"hs_shift":shift}),"A count discrepancy must not create a balancing entry")

    def test_readable_presentation_hides_internal_keys_and_keeps_explicit_trace_view(self):
        entered=self.enter()
        charge=frappe.db.get_value("HS Charge",{"stay":entered["stay"]},"name")
        view=api.report_page(self.prop.name,"Producción",dict(room=self.room.name))
        self.assertNotIn("name",view["columns"])
        self.assertNotIn("stay",view["columns"])
        self.assertNotIn("folio",view["columns"])
        row=next(r for r in view["rows"] if r.name==charge)
        self.assertEqual(row._display["item"],frappe.db.get_value("Item",row.item,"item_name"))
        self.assertEqual(row._display["room"],self.room.room_number)
        self.assertEqual(row.stay,entered["stay"],"Readability must not rename or replace command keys")
        trace=api.report_page(self.prop.name,"Producción",dict(room=self.room.name,detail_level="technical"))
        self.assertIn("name",trace["columns"])
        self.assertEqual(next(r.name for r in trace["rows"] if r.name==charge),charge)
        frappe.local.response=frappe._dict()
        api.export_report(self.prop.name,"Producción",dict(room=self.room.name),"xlsx")
        wb=load_workbook(io.BytesIO(frappe.local.response.filecontent),read_only=True)
        values=list(wb.active.values)
        flat={str(v) for r in values for v in r if v is not None}
        self.assertNotIn(charge,flat)
        self.assertNotIn(entered["stay"],flat)
        self.assertNotIn(row.item,flat)
        self.assertIn(row._display["item"],flat)
        wb.close()

    def test_availability_shares_capacity_with_roomless_reservations(self):
        start=now_datetime().replace(microsecond=0)+timedelta(days=20)
        end=start+timedelta(days=1)
        values=dict(starts_at=str(start),ends_at=str(end),guest=self.guest.name,customer=self.guest.customer,
                    rooms=[dict(category=self.room.category,people=1)])
        values["acceptance"]=api.preview(self.prop.name,values)["acceptance"]
        self.call("reserve",values)
        filters=dict(starts_at=str(start),ends_at=str(end),category=self.room.category)
        rows=api.report(self.prop.name,"Disponibilidad",filters)
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0].available_capacity,rows[0].total_rooms-1)
        self.assertEqual(len(rows[0].candidate_rooms.split(", ")),rows[0].total_rooms)
        from hospitality_suite.availability import check,capacity
        self.assertEqual(len(check(self.prop,self.room.category,start,end)),rows[0].total_rooms)
        proposed=[frappe._dict(category=self.room.category,room=None,starts_at=start,ends_at=end,buffer_minutes=0) for _ in range(rows[0].total_rooms-1)]
        self.assertEqual(capacity(self.prop,self.room.category,start,end,proposed=proposed)["available"],0)
        with self.assertRaises(frappe.ValidationError):
            check(self.prop,self.room.category,start,end,proposed=proposed)

    def test_company_scope_only_aggregates_authorized_properties_and_options(self):
        company=frappe.get_doc(dict(doctype="Company",company_name="Report isolation "+uuid.uuid4().hex[:7],abbr=uuid.uuid4().hex[:5].upper(),country="Nicaragua",default_currency="NIO",chart_of_accounts="Standard")).insert()
        with operation():
            foreign=frappe.get_doc(dict(doctype="HS Property",property_name="Report foreign fixture",company=company.name,mode="Hotel",time_zone="Asia/Tokyo",currency="NIO",business_date=frappe.utils.nowdate(),financial_enabled=0)).insert()
        user="report-"+uuid.uuid4().hex[:8]+"@example.invalid"
        frappe.get_doc(dict(doctype="User",email=user,first_name="Report fixture",send_welcome_email=0,user_type="System User",roles=[dict(role="HS Manager")])).insert()
        with operation():
            frappe.get_doc(dict(doctype="HS Access",company=self.prop.company,property=self.prop.name,user=user)).insert()
        frappe.set_user(user)
        rows=api.report("","Disponibilidad",dict(company=self.prop.company))
        self.assertTrue(rows)
        self.assertEqual({r.property for r in rows},{self.prop.name})
        options=api.report_options("","Producción",self.prop.company)
        self.assertEqual({r.name for r in options["properties"]},{self.prop.name})
        self.assertTrue(all(r.property_label==self.prop.property_name for r in options["rooms"]))
        with self.assertRaises(frappe.PermissionError):
            api.report("","Auditoría",dict(company=company.name))
        with self.assertRaises(frappe.PermissionError):
            api.report(self.prop.name,"Auditoría",dict(company=company.name))
        with self.assertRaises(frappe.PermissionError):
            api.report_options(foreign.name,"Producción")
        with self.assertRaises(frappe.PermissionError):
            api.export_report("","Producción",dict(company=company.name))
        frappe.set_user("housekeeping@hospitality.example.invalid")
        with self.assertRaises(frappe.PermissionError):
            api.report_options(self.prop.name,"Producción")

    def test_only_positive_pending_folios_and_applicable_report_options(self):
        entered=self.enter()
        rows=api.report(self.prop.name,"Cuentas pendientes",dict(room=self.room.name))
        self.assertEqual([r.name for r in rows],[entered["folio"]])
        self.assertGreater(decimal(rows[0].outstanding),0)
        with operation():
            frappe.get_doc(dict(doctype="HS Folio",property=self.prop.name,company=self.prop.company,stay=entered["stay"],customer=self.guest.customer,currency=self.prop.currency,status="Open",label="Cuenta vacía de prueba")).insert()
        self.assertEqual(len(api.report(self.prop.name,"Cuentas pendientes",dict(room=self.room.name))),1)
        from hospitality_suite.reports import REPORTS
        for name in REPORTS:
            options=api.report_options(self.prop.name,name)
            self.assertTrue(options["date_fields"])
            self.assertIsInstance(api.report(self.prop.name,name),list)
        self.assertNotIn("room",api.report_options(self.prop.name,"Auditoría")["filters"])
