from datetime import timedelta
import uuid

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import now_datetime

from hospitality_suite import api
from hospitality_suite.service import operation


class TestFinance(IntegrationTestCase):
    def setUp(self):
        super().setUp()
        self.addCleanup(frappe.db.rollback)
        frappe.set_user("Administrator")
        self.prop = frappe.get_doc("HS Property", {"property_name":"Hotel Jardín Demo"})
        self.room = frappe.get_all("HS Room", filters={"property":self.prop.name}, fields=["*"])[0]
        self.guest = frappe.get_all("HS Guest", filters={"property":self.prop.name}, fields=["*"])[0]
        start = now_datetime().replace(microsecond=0)
        data = dict(starts_at=str(start), ends_at=str(start+timedelta(days=1)), guest=self.guest.name,
                    customer=self.guest.customer, rooms=[dict(category=self.room.category, room=self.room.name, people=1)])
        data["acceptance"] = api.preview(self.prop.name, data)["acceptance"]
        self.stay = self.run_command("walk_in", data)
        self.shift = self.run_command("open_shift", dict(opening_amount=0))["shift"]
        self.method = frappe.db.get_value("HS Payment Method", {"property":self.prop.name})

    def tearDown(self):
        frappe.set_user("Administrator")
        super().tearDown()

    def run_command(self, action, data, key=None):
        return api.execute(self.prop.name, action, key or uuid.uuid4().hex, data)

    def pay(self, amount, kind, **extra):
        return self.run_command("payment", dict(stay=self.stay["stay"], folio=self.stay["folio"], shift=self.shift,
                                               method=self.method, amount=amount, kind=kind, **extra))

    def test_guarantee_is_liability_then_refund(self):
        guarantee = self.pay(200, "Guarantee")
        self.assertIsNotNone(guarantee["journal_entry"])
        journal = frappe.get_doc("Journal Entry", guarantee["journal_entry"])
        self.assertEqual(sum(a.credit_in_account_currency for a in journal.accounts if a.account == self.prop.guarantee_account), 200)
        self.pay(200, "Guarantee Refund", source_payment=guarantee["payment"], reason="Salida sin daños")

    def test_guarantee_application_is_standard_advance(self):
        guarantee = self.pay(200, "Guarantee")
        self.run_command("apply_guarantee", dict(source_payment=guarantee["payment"], folio=self.stay["folio"], amount=200, reason="Aplicación autorizada al alojamiento"))
        result = self.run_command("invoice", dict(folio=self.stay["folio"]))
        inv = frappe.get_doc("Sales Invoice", result["invoice"])
        self.assertEqual(inv.total_advance, 200)
        self.assertEqual(inv.outstanding_amount, 1300)

    def test_double_payment_is_one_posting(self):
        key = uuid.uuid4().hex
        values = dict(stay=self.stay["stay"], folio=self.stay["folio"], shift=self.shift, method=self.method, amount=100, kind="Payment")
        one = self.run_command("payment", values, key)
        two = self.run_command("payment", values, key)
        self.assertEqual(one, two)
        self.assertEqual(frappe.db.count("HS Payment", {"payment_entry":one["payment_entry"]}), 1)

    def test_close_difference_needs_supervisor_and_blocks_new_payment(self):
        self.pay(100, "Advance")
        shift = frappe.get_doc("HS Shift", self.shift)
        result = self.run_command("close_shift", dict(shift=shift.name, modified=str(shift.modified), denominations={"NIO":[dict(value=1,count=90)]}, reason="Faltante de prueba"))
        self.assertEqual(result["status"], "Review")
        with self.assertRaises(frappe.ValidationError):
            self.pay(10,"Advance")

    def test_multicurrency_advance_preserves_exchange(self):
        company = frappe.get_doc("Company", self.prop.company)
        bank = frappe.get_doc(dict(doctype="Account", account_name="USD Caja Test " + uuid.uuid4().hex[:6], company=company.name,
                                  parent_account=frappe.db.get_value("Account", {"company":company.name,"account_type":"Cash","is_group":1}),
                                  account_type="Cash", account_currency="USD")).insert()
        with operation():
            method = frappe.get_doc(dict(doctype="HS Payment Method", property=self.prop.name, company=company.name,
                                        account=bank.name, currency="USD", mode_of_payment="Efectivo demo", is_cash=1)).insert()
        result = self.run_command("payment", dict(stay=self.stay["stay"], folio=self.stay["folio"], shift=self.shift,
                                                 method=method.name, amount=10, kind="Advance"))
        record = frappe.get_doc("HS Payment", result["payment"])
        self.assertEqual(record.currency,"USD")
        self.assertAlmostEqual(record.exchange_rate,36.6243)
        self.assertEqual(record.base_amount,366.24)
        settled=self.pay(1133.76,"Payment")
        self.assertTrue(settled["payment_entry"])
        self.assertEqual(api.stay_detail(self.stay["stay"])["folios"][0]["balance"]["outstanding"],"0.00")

    def test_standard_payment_cancellation_updates_folio(self):
        result = self.pay(1500,"Payment")
        pe = frappe.get_doc("Payment Entry", result["payment_entry"])
        pe.hs_correction_reason="Prueba de anulación desde ERPNext"
        pe.cancel()
        self.assertEqual(frappe.db.get_value("HS Payment", result["payment"],"status"), "Cancelled")
        detail = api.stay_detail(self.stay["stay"])
        self.assertEqual(detail["folios"][0]["balance"]["outstanding"],"1500.00")

    def test_partial_advance_refund_then_invoice(self):
        original = self.pay(500,"Advance")
        self.pay(100,"Refund",source_payment=original["payment"],reason="Devolución parcial autorizada")
        result = self.run_command("invoice",dict(folio=self.stay["folio"]))
        inv = frappe.get_doc("Sales Invoice",result["invoice"])
        self.assertEqual(inv.total_advance,400)
        self.assertEqual(inv.outstanding_amount,1100)

    def test_credit_note_uses_standard_return(self):
        result = self.run_command("invoice",dict(folio=self.stay["folio"]))
        credit = self.run_command("credit_note",dict(invoice=result["invoice"],reason="Anulación autorizada de prueba"))
        inv = frappe.get_doc("Sales Invoice",credit["credit_note"])
        self.assertEqual(inv.is_return,1)
        self.assertEqual(inv.return_against,result["invoice"])
        self.assertEqual(inv.grand_total,-1500)

    def test_credit_note_and_paid_invoice_refund(self):
        paid=self.pay(1500,"Payment")
        original=frappe.get_doc("Payment Entry",paid["payment_entry"]).references[0].reference_name
        credit=self.run_command("credit_note",dict(invoice=original,reason="Devolución de alojamiento autorizada"))
        result=self.pay(1500,"Refund",source_payment=paid["payment"],credit_note=credit["credit_note"],reason="Devolución autorizada")
        self.assertTrue(result["payment_entry"])
        self.assertEqual(frappe.db.get_value("Sales Invoice",credit["credit_note"],"outstanding_amount"),0)

    def test_advance_after_invoice_from_closed_shift_reconciles_without_cash_change(self):
        invoice=self.run_command("invoice",dict(folio=self.stay["folio"]))["invoice"]
        paid=self.pay(500,"Advance")
        shift=frappe.get_doc("HS Shift",self.shift)
        self.run_command("close_shift",dict(shift=shift.name,modified=str(shift.modified),denominations={"NIO":[dict(value=100,count=5)]}))
        original=frappe.get_doc("Payment Entry",paid["payment_entry"])
        result=self.run_command("invoice",dict(folio=self.stay["folio"]))
        self.assertEqual(result["reconciled"][0]["amount"],500)
        original.reload()
        self.assertEqual(original.unallocated_amount,0)
        self.assertEqual(original.received_amount,500)
        self.assertEqual(original.hs_shift,self.shift)
        self.assertEqual(frappe.db.get_value("Sales Invoice",invoice,"outstanding_amount"),1000)
        from hospitality_suite.cashier import expected
        self.assertEqual(expected(shift)["NIO"],"500.00")

    def test_guarantee_applied_to_existing_invoice_is_not_cash_twice(self):
        guarantee=self.pay(200,"Guarantee")
        invoice=self.run_command("invoice",dict(folio=self.stay["folio"]))["invoice"]
        self.run_command("apply_guarantee",dict(source_payment=guarantee["payment"],folio=self.stay["folio"],amount=200,reason="Aplicación autorizada"))
        self.assertEqual(frappe.db.get_value("Sales Invoice",invoice,"outstanding_amount"),1300)
        from hospitality_suite.cashier import expected
        self.assertEqual(expected(frappe.get_doc("HS Shift",self.shift))["NIO"],"200.00")

    def test_cash_change_and_fractional_count(self):
        result=self.pay(95,"Advance",tendered=100)
        self.assertEqual(result["change"],"5.00")
        from hospitality_suite.cashier import expected
        self.assertEqual(expected(frappe.get_doc("HS Shift",self.shift))["NIO"],"95.00")
        shift=frappe.get_doc("HS Shift",self.shift)
        with self.assertRaises(frappe.ValidationError):
            self.run_command("close_shift",dict(shift=self.shift,modified=str(shift.modified),denominations={"NIO":[dict(value=100,count=.95)]}))

    def test_invoice_and_payment_created_outside_pms_sync_and_reverse(self):
        source=frappe.get_doc("HS Charge",{"stay":self.stay["stay"],"status":"Pending"})
        invoice=frappe.get_doc(dict(doctype="Sales Invoice",company=self.prop.company,customer=self.guest.customer,currency=self.prop.currency,
                                   debit_to=self.prop.receivable_account,cost_center=self.prop.cost_center,hs_property=self.prop.name,
                                   hs_stay=self.stay["stay"],hs_folio=self.stay["folio"],update_stock=0,
                                   items=[dict(item_code=source.item,qty=source.quantity,rate=source.rate,cost_center=self.prop.cost_center,hs_charge=source.name)]))
        invoice.insert().submit()  # Deliberately bypasses PMS posting helper.
        source.reload()
        self.assertEqual(source.sales_invoice,invoice.name)
        self.assertEqual(source.status,"Invoiced")
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        payment=get_payment_entry("Sales Invoice",invoice.name,bank_account=frappe.db.get_value("HS Payment Method",self.method,"account"))
        payment.update(dict(hs_property=self.prop.name,hs_stay=self.stay["stay"],hs_folio=self.stay["folio"],hs_shift=self.shift,
                            mode_of_payment=frappe.db.get_value("HS Payment Method",self.method,"mode_of_payment"),reference_no="EXTERNAL-TEST",reference_date=frappe.utils.nowdate()))
        payment.insert().submit()
        record=frappe.get_doc("HS Payment",{"payment_entry":payment.name})
        self.assertEqual(record.shift,self.shift)
        self.assertEqual(record.amount,1500)
        self.assertEqual(api.stay_detail(self.stay["stay"])["folios"][0]["balance"]["outstanding"],"0.00")
        payment.hs_correction_reason="Reversión externa de prueba"
        payment.cancel()
        invoice.reload()
        invoice.hs_correction_reason="Reversión externa de prueba"
        invoice.cancel()
        self.assertEqual(frappe.db.get_value("HS Payment",record.name,"status"),"Cancelled")
        self.assertEqual(frappe.db.get_value("HS Charge",source.name,"status"),"Pending")

    def test_native_notifications_never_use_global_list_channel(self):
        from unittest.mock import patch
        result=self.run_command("invoice",dict(folio=self.stay["folio"]))
        invoice=frappe.get_doc("Sales Invoice",result["invoice"])
        with patch("frappe.publish_realtime") as publish:
            invoice.notify_update()
        listed=[c for c in publish.call_args_list if c.args[0]=="list_update"]
        self.assertTrue(listed)
        self.assertTrue(all(c.kwargs.get("user") for c in listed))
        self.assertNotIn("housekeeping@hospitality.example.invalid",[c.kwargs["user"] for c in listed])

    def external_invoice(self,source,sources,rate):
        return frappe.get_doc(dict(doctype="Sales Invoice",company=self.prop.company,customer=self.guest.customer,currency=self.prop.currency,
                                  debit_to=self.prop.receivable_account,cost_center=self.prop.cost_center,hs_property=self.prop.name,
                                  hs_stay=self.stay["stay"],hs_folio=self.stay["folio"],update_stock=0,
                                  items=[dict(item_code=source.item,qty=source.quantity,rate=rate,price_list_rate=source.rate,
                                              cost_center=self.prop.cost_center,hs_charge=source.name,hs_charge_sources=frappe.as_json(sources))]))

    def test_external_invoice_requires_all_adjustments_and_syncs_every_source(self):
        source=frappe.get_doc("HS Charge",{"stay":self.stay["stay"],"status":"Pending"})
        adjustment=self.run_command("adjustment",dict(charge=source.name,adjustment_type="Discount",amount=100,reason="Descuento ficticio autorizado"))["charge"]
        with self.assertRaises(frappe.ValidationError):
            self.external_invoice(source,[source.name],source.rate).insert()
        invoice=self.external_invoice(source,[source.name,adjustment],1400).insert()
        invoice.submit()
        for name in (source.name,adjustment):
            record=frappe.get_doc("HS Charge",name)
            self.assertEqual((record.status,record.sales_invoice,record.invoice_item),("Invoiced",invoice.name,invoice.items[0].name))
        from hospitality_suite.reconciliation import differences
        self.assertEqual(differences(self.prop),[])
        invoice.hs_correction_reason="Anulación de prueba"
        invoice.cancel()
        self.assertEqual(frappe.db.get_value("HS Charge",adjustment,"status"),"Pending")

    def test_external_invoice_rejects_reversed_charge(self):
        source=frappe.get_doc("HS Charge",{"stay":self.stay["stay"],"status":"Pending"})
        with operation():
            source.status="Reversed"
            source.save()
        with self.assertRaises(frappe.ValidationError):
            self.external_invoice(source,[source.name],source.rate).insert()

    def test_used_payment_method_cannot_reclassify_history(self):
        self.pay(100,"Advance")
        method=frappe.get_doc("HS Payment Method",self.method)
        method.is_cash=not method.is_cash
        with self.assertRaises(frappe.ValidationError):
            method.save()

    def test_native_currency_precision_three_and_zero(self):
        from hospitality_suite.pricing import money
        original=frappe.db.get_default("currency_precision")
        try:
            frappe.db.set_default("currency_precision","3")
            self.assertEqual(str(money("10.1235")),"10.124")
            result=self.pay("100.123","Advance")
            self.assertEqual(frappe.db.get_value("HS Payment",result["payment"],"amount"),100.123)
            self.assertEqual(frappe.db.get_value("Payment Entry",result["payment_entry"],"received_amount"),100.123)
            from hospitality_suite.cashier import expected
            self.assertEqual(expected(frappe.get_doc("HS Shift",self.shift))["NIO"],"100.123")
            import io
            from openpyxl import load_workbook
            records=api.report(self.prop.name,"Cobros")
            self.assertEqual(next(r.base_currency for r in records if r.name==result["payment"]),"NIO")
            def check_export(number_format):
                api.export_report(self.prop.name,"Cobros",file_format="xlsx")
                book=load_workbook(io.BytesIO(frappe.local.response.filecontent),read_only=True)
                sheet=book.active
                headers=[cell.value for cell in next(sheet.rows)]
                column=headers.index("Importe")+1
                cell=sheet.cell(2,column)
                self.assertEqual(cell.value,100.123)
                self.assertEqual(cell.number_format,number_format)
                book.close()
            check_export("#,##0.000")
            frappe.db.set_default("currency_precision","0")
            check_export("#,##0")
            self.assertEqual(str(money("10.5")),"11")
            start=now_datetime().replace(microsecond=0)
            values=dict(starts_at=str(start),ends_at=str(start+timedelta(days=1)),customer=self.guest.customer,guest=self.guest.name,
                        rooms=[dict(category=self.room.category,people=1)])
            quote=api.preview(self.prop.name,values)
            self.assertEqual(quote["lines"][0]["precision"],0)
            self.assertEqual(quote["amount"],"1500")
        finally:
            frappe.db.set_default("currency_precision",original)

    def test_untracked_native_cash_is_visible_and_blocks_shift_close(self):
        journal=frappe.get_doc(dict(doctype="Journal Entry",company=self.prop.company,posting_date=frappe.utils.nowdate(),
                                   hs_property=self.prop.name,hs_shift=self.shift,
                                   accounts=[dict(account=frappe.db.get_value("HS Payment Method",self.method,"account"),debit_in_account_currency=50),
                                             dict(account=self.prop.guarantee_account,credit_in_account_currency=50)] )).insert()
        journal.submit()
        from hospitality_suite.reconciliation import differences
        self.assertTrue(any(r["erp_document"]==journal.name for r in differences(self.prop)))
        shift=frappe.get_doc("HS Shift",self.shift)
        with self.assertRaises(frappe.ValidationError):
            self.run_command("close_shift",dict(shift=shift.name,modified=str(shift.modified),denominations={}))
        journal.hs_correction_reason="Anulación de movimiento sin conciliar"
        journal.cancel()
        self.assertEqual(differences(self.prop),[])

    def test_departed_visit_cannot_receive_new_advance(self):
        self.pay(1500,"Payment")
        stay=frappe.get_doc("HS Stay",self.stay["stay"])
        self.run_command("check_out",dict(stay=stay.name,modified=str(stay.modified)))
        with self.assertRaises(frappe.ValidationError):
            self.pay(100,"Advance")

    def test_withdrawal_posts_transfer_once_and_reduces_expected_cash(self):
        parent=frappe.db.get_value("Account",{"company":self.prop.company,"account_type":"Cash","is_group":1},"name")
        account=frappe.get_doc(dict(doctype="Account",company=self.prop.company,parent_account=parent,account_name="Entrega ficticia "+uuid.uuid4().hex[:6],account_type="Cash",account_currency="NIO")).insert()
        with operation():
            self.prop.cash_transfer_account=account.name
            self.prop.save()
        self.pay(500,"Advance")
        key=uuid.uuid4().hex
        values=dict(shift=self.shift,method=self.method,amount=200,reason="Entrega a supervisión ficticia")
        result=self.run_command("withdrawal",values,key)
        self.assertEqual(result,self.run_command("withdrawal",values,key))
        from hospitality_suite.cashier import expected
        self.assertEqual(expected(frappe.get_doc("HS Shift",self.shift))["NIO"],"300.00")
        journal=frappe.get_doc("Journal Entry",result["journal_entry"])
        self.assertEqual(sum(r.debit_in_account_currency for r in journal.accounts if r.account==account.name),200)
        with operation():
            self.prop.cash_transfer_account=frappe.db.get_value("HS Payment Method",self.method,"account")
            self.prop.save()
        with self.assertRaises(frappe.ValidationError):
            self.run_command("withdrawal",values)

    def test_usd_guarantee_refund_preserves_liability_and_posts_exchange_difference(self):
        from decimal import Decimal
        from unittest.mock import patch
        parent=frappe.db.get_value("Account",{"company":self.prop.company,"account_type":"Cash","is_group":1},"name")
        cash=frappe.get_doc(dict(doctype="Account",company=self.prop.company,parent_account=parent,account_name="USD garantía ficticia "+uuid.uuid4().hex[:6],account_type="Cash",account_currency="USD")).insert()
        with operation():
            method=frappe.get_doc(dict(doctype="HS Payment Method",property=self.prop.name,company=self.prop.company,account=cash.name,currency="USD",mode_of_payment="Efectivo demo",is_cash=1)).insert()
        original=self.run_command("payment",dict(stay=self.stay["stay"],folio=self.stay["folio"],shift=self.shift,method=method.name,amount=10,kind="Guarantee"))
        with patch("hospitality_suite.accounting.exchange",side_effect=lambda source,target,*args: Decimal(37) if source=="USD" else Decimal(1)):
            returned=self.run_command("payment",dict(stay=self.stay["stay"],shift=self.shift,method=method.name,amount=10,kind="Guarantee Refund",source_payment=original["payment"],reason="Devolución con cambio de referencia ficticio"))
        entries=frappe.get_all("GL Entry",filters={"voucher_type":"Journal Entry","voucher_no":["in",[original["journal_entry"],returned["journal_entry"]]],"account":self.prop.guarantee_account},fields=["debit","credit"])
        self.assertAlmostEqual(sum(r.debit-r.credit for r in entries),0,places=2)
        journal=frappe.get_doc("Journal Entry",returned["journal_entry"])
        difference=frappe.db.get_value("Company",self.prop.company,"exchange_gain_loss_account")
        self.assertAlmostEqual(sum(r.debit-r.credit for r in journal.accounts if r.account==difference),3.76,places=2)
